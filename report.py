"""Excel export: 3 sheets, item x (tier.enchant) matrix with merged tier
headers (ТЗ §9)."""
import os
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import storage
from config_io import load_config

CLASS_SHEETS = [
    ("hunter", "Охотник"),
    ("mage", "Маг"),
    ("warrior", "Воин"),
]

HEADER_FONT = Font(bold=True)
CENTER = Alignment(horizontal="center", vertical="center")
REVIEW_FONT = Font(italic=True, color="CC0000")
HEADER_BORDER = Border(*(Side(style="thin", color="999999"),) * 4)


def _fill(hex_color: str) -> PatternFill:
    return PatternFill(start_color=hex_color, end_color=hex_color, fill_type="solid")


# Fixed report template: always render tiers 4-8 x enchants 0-3, regardless
# of what's actually in the DB, so the sheet reads the same every export.
FIXED_TIERS = list(range(4, 9))
FIXED_ENCHANTS = list(range(0, 4))

# Tier / enchant colors match Albion Online's own item-quality palette.
TIER_FILLS = {
    4: _fill("ADD8E6"),  # голубой
    5: _fill("FF6961"),  # красный
    6: _fill("FFA500"),  # оранжевый
    7: _fill("FFFF00"),  # желтый
    8: _fill("FFFFFF"),  # белый (с рамкой, иначе теряется на белом фоне)
}
ENCHANT_FILLS = {
    0: _fill("D9D9D9"),  # серый
    1: _fill("92D050"),  # зеленый
    2: _fill("00B0F0"),  # синий
    3: _fill("7030A0"),  # фиолетовый
}
ITEM_HEADER_FILL = _fill("D9D9D9")  # серый, как и зачарование 0


def build_sheet(ws, rows: list) -> None:
    price_lookup = {}
    review_lookup = {}
    for r in rows:
        key = (r["item_name"], r["tier"], r["enchant"])
        price_lookup[key] = r["price"]
        review_lookup[key] = r["needs_review"]

    items = sorted({
        r["item_name"] for r in rows
        if r["tier"] in FIXED_TIERS and r["enchant"] in FIXED_ENCHANTS
    })

    combo_to_col = {}
    tier_col_span = {}
    col = 2  # column A is the item name
    for tier in FIXED_TIERS:
        start_col = col
        for enchant in FIXED_ENCHANTS:
            combo_to_col[(tier, enchant)] = col
            col += 1
        tier_col_span[tier] = (start_col, col - 1)
    last_col = col - 1

    header_cell = ws.cell(row=1, column=1, value="Предмет")
    header_cell.font = HEADER_FONT
    header_cell.fill = ITEM_HEADER_FILL
    header_cell.alignment = CENTER
    header_cell.border = HEADER_BORDER
    ws.merge_cells(start_row=1, start_column=1, end_row=2, end_column=1)
    ws.cell(row=2, column=1).border = HEADER_BORDER

    for tier in FIXED_TIERS:
        start_col, end_col = tier_col_span[tier]
        ws.merge_cells(start_row=1, start_column=start_col, end_row=1, end_column=end_col)
        for c in range(start_col, end_col + 1):
            tcell = ws.cell(row=1, column=c)
            tcell.fill = TIER_FILLS[tier]
            tcell.border = HEADER_BORDER
        cell = ws.cell(row=1, column=start_col, value=f"Тир {tier}")
        cell.font = HEADER_FONT
        cell.alignment = CENTER

        for enchant in FIXED_ENCHANTS:
            c = combo_to_col[(tier, enchant)]
            ecell = ws.cell(row=2, column=c, value=enchant)
            ecell.font = HEADER_FONT
            ecell.alignment = CENTER
            ecell.fill = ENCHANT_FILLS[enchant]
            ecell.border = HEADER_BORDER

    for i, item_name in enumerate(items):
        row_idx = 3 + i
        ws.cell(row=row_idx, column=1, value=item_name)
        for tier in FIXED_TIERS:
            for enchant in FIXED_ENCHANTS:
                c = combo_to_col[(tier, enchant)]
                key = (item_name, tier, enchant)
                price = price_lookup.get(key)
                cell = ws.cell(row=row_idx, column=c, value=price if price else "—")
                if review_lookup.get(key):
                    cell.font = REVIEW_FONT

    ws.freeze_panes = "B3"
    ws.column_dimensions["A"].width = 30
    for c in range(2, last_col + 1):
        ws.column_dimensions[get_column_letter(c)].width = 10

    # Captures that don't fit the fixed template (unresolved alias -> tier
    # outside 4-8, or an enchant level outside 0-3) aren't dropped
    # silently; list them below the matrix.
    leftover = sorted(
        (r for r in rows if r["tier"] not in FIXED_TIERS or r["enchant"] not in FIXED_ENCHANTS),
        key=lambda r: r["item_name"],
    )
    if leftover:
        start_row = 3 + len(items)
        note_cell = ws.cell(
            row=start_row, column=1,
            value="Требует проверки (тир вне 4-8 или зачарование вне 0-3)",
        )
        note_cell.font = HEADER_FONT
        for j, title in enumerate(["Предмет", "Тир", "Зачарование", "Цена"]):
            ws.cell(row=start_row + 1, column=1 + j, value=title).font = HEADER_FONT
        for i, r in enumerate(leftover):
            row_idx = start_row + 2 + i
            tier_display = r["tier"] if r["tier"] >= 0 else "?"
            ws.cell(row=row_idx, column=1, value=r["item_name"])
            ws.cell(row=row_idx, column=2, value=tier_display)
            ws.cell(row=row_idx, column=3, value=r["enchant"])
            ws.cell(row=row_idx, column=4, value=r["price"] if r["price"] else "—")


def generate_report(output_path: str = None) -> str:
    config = load_config()
    db_path = storage.resolve_db_path(config["db_path"])
    storage.init_db(db_path)

    wb = Workbook()
    wb.remove(wb.active)

    for class_key, sheet_title in CLASS_SHEETS:
        ws = wb.create_sheet(title=sheet_title)
        rows = storage.fetch_by_class(db_path, class_key)
        build_sheet(ws, rows)

    if output_path is None:
        out_dir = os.path.expandvars(config.get("report_output_dir", "./reports"))
        os.makedirs(out_dir, exist_ok=True)
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        output_path = os.path.join(out_dir, f"albion_report_{timestamp}.xlsx")
    else:
        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)

    wb.save(output_path)
    return output_path
